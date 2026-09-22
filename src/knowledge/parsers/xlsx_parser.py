"""Bounded-memory Excel reader. Only one row/block is resident at a time.

openpyxl's reader adapter preserves cached formula values, dates and rich text.
Its shared-string table is disk-backed; sparse rows are spooled per sheet so
format-only cells and globally empty columns never expand the Markdown output.
The two openpyxl reader internals used here are covered by compatibility tests.
"""
from __future__ import annotations

import json
import sqlite3
import tempfile
from functools import lru_cache
from pathlib import Path
from typing import Callable, Iterator
from xml.etree.ElementTree import iterparse

from openpyxl.cell.text import Text
from openpyxl.reader.excel import ExcelReader
from openpyxl.worksheet._reader import WorkSheetParser
from openpyxl.xml.constants import SHARED_STRINGS, SHEET_MAIN_NS
from openpyxl.utils import get_column_letter

from src.knowledge.parsers.base import ParsedDocument, ParsedSection

BLOCK_CHARS = 32768
Progress = Callable[[str, dict], None]


class _DiskStrings:
    def __init__(self, path: Path):
        self.db = sqlite3.connect(path)
        self.db.execute("PRAGMA cache_size=-1024")
        self.db.execute("CREATE TABLE strings (id INTEGER PRIMARY KEY, value TEXT)")
        self.get = lru_cache(maxsize=128)(self._get)

    def _get(self, index):
        row = self.db.execute("SELECT value FROM strings WHERE id=?", (index,)).fetchone()
        if row is None:
            raise IndexError(index)
        return row[0]

    def __getitem__(self, index):
        return self.get(index)

    def close(self):
        self.get.cache_clear()
        self.db.close()


class _StreamingReader(ExcelReader):
    def __init__(self, path, strings, progress):
        self.disk_strings = strings
        self.progress = progress
        super().__init__(path, read_only=True, data_only=True, keep_links=False)

    def read_strings(self):
        self.shared_strings = self.disk_strings
        part = self.package.find(SHARED_STRINGS)
        if part is None:
            return
        tag = "{%s}si" % SHEET_MAIN_NS
        with self.archive.open(part.PartName.lstrip('/')) as source:
            events = iterparse(source, events=("start", "end"))
            _, root = next(events)
            count = 0
            for event, node in events:
                if event == "end" and node.tag == tag:
                    value = Text.from_tree(node).content.replace('x005F_', '')
                    self.disk_strings.db.execute("INSERT INTO strings VALUES (?, ?)", (count, value))
                    count += 1
                    root.clear()
                    if self.progress and count % 2048 == 0:
                        self.progress("parsing", {"detail": f"读取共享文本：{count} 条"})
            self.disk_strings.db.commit()


def _row_to_md(cells) -> str:
    return "| " + " | ".join(
        ("" if c is None else str(c)).replace("|", "\\|").replace("\n", " ").replace("\r", " ")
        for c in cells
    ) + " |"


def iter_xlsx_sections(path: Path, progress: Progress | None = None) -> Iterator[ParsedSection]:
    """Yield bounded blocks with original sheet/row coordinates; always close resources."""
    with tempfile.TemporaryDirectory(prefix="insight-xlsx-") as temp:
        strings = _DiskStrings(Path(temp) / "strings.db")
        reader = None
        try:
            if progress:
                progress("parsing", {"detail": "读取工作簿"})
            reader = _StreamingReader(path, strings, progress)
            reader.read()
            wb = reader.wb
            section_index = 0
            for sheet_index, ws in enumerate(wb.worksheets):
                prefix = f"工作表 {sheet_index + 1}/{len(wb.worksheets)}：{ws.title}"
                if progress:
                    progress("parsing", {"detail": prefix})
                # Disk spool permits column filtering without retaining the sheet.
                with tempfile.TemporaryFile(mode="w+", encoding="utf-8") as spool:
                    columns = set()
                    row_count = 0
                    with ws._get_source() as source:
                        parser = WorkSheetParser(source, strings, data_only=True,
                            epoch=wb.epoch, date_formats=wb._date_formats,
                            timedelta_formats=wb._timedelta_formats)
                        for row_index, row in parser.parse():
                            values = {c["column"]: str(c["value"]) for c in row
                                      if c["value"] is not None and str(c["value"]).strip()}
                            if values:
                                columns.update(values)
                                spool.write(json.dumps([row_index, values], ensure_ascii=False) + "\n")
                                row_count += 1
                            if progress and row_index % 512 == 0:
                                progress("parsing", {"detail": f"{prefix} · 已读取 {row_index} 行"})
                    spool.seek(0)
                    columns = sorted(columns)
                    labels = ",".join(get_column_letter(c) for c in columns)
                    header = None
                    lines = []
                    size = 0
                    start = end = 0
                    def section():
                        return ParsedSection(
                            text="\n".join(lines), source_path=path, section_index=section_index,
                            section_label=f"Sheet: {ws.title} · 行 {start}-{end}",
                            extra={"sheet": ws.title, "row_start": start, "row_end": end,
                                   "row_count": row_count, "columns": labels})
                    for encoded in spool:
                        row_index, values = json.loads(encoded)
                        line = _row_to_md(values.get(str(c)) for c in columns)
                        if header is None:
                            header = [line, "| " + " | ".join("---" for _ in columns) + " |"]
                            lines = header.copy()
                            size = sum(map(len, lines))
                            start = end = row_index
                            continue
                        if size + len(line) > BLOCK_CHARS and end > start:
                            yield section()
                            section_index += 1
                            lines = header.copy()
                            size = sum(map(len, lines))
                            start = row_index
                        lines.append(line)
                        size += len(line) + 1
                        end = row_index
                    if lines:
                        yield section()
                        section_index += 1
                if progress:
                    progress("parsing", {"detail": f"{prefix} · {row_count} 行已解析并切片",
                                         "completed": sheet_index + 1, "total": len(wb.worksheets)})
        finally:
            if reader is not None:
                reader.archive.close()
            strings.close()


def parse_xlsx(path: Path) -> ParsedDocument:
    # Compatibility API for callers that explicitly need a materialized document.
    sections = list(iter_xlsx_sections(path))
    return ParsedDocument(source_path=path, sections=sections, file_type="xlsx",
                          title=path.stem,
                          metadata={"sheet_count": len({s.extra["sheet"] for s in sections})})
