"""PowerPoint .pptx 解析：python-pptx。

策略：每张幻灯片一个 section，包含标题 + 正文 + 表格。
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from src.knowledge.parsers.base import ParsedDocument, ParsedSection


def parse_pptx(path: Path) -> ParsedDocument:
    """解析 .pptx，每张幻灯片一个 section。"""
    prs = Presentation(str(path))
    sections: list[ParsedSection] = []
    for i, slide in enumerate(prs.slides):
        lines: list[str] = []
        title_text = ""
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text.strip()
                if not text:
                    continue
                if not title_text and shape == slide.shapes.title:
                    title_text = text
                else:
                    lines.append(text)
            if shape.has_table:
                tbl = shape.table
                for row in tbl.rows:
                    cells = [(c.text or "").strip() for c in row.cells]
                    lines.append(" | ".join(cells))
        label = f"Slide {i + 1}"
        if title_text:
            label += f": {title_text[:50]}"
        sections.append(
            ParsedSection(
                text="\n".join(lines),
                source_path=path,
                section_index=i,
                section_label=label,
                extra={"slide": i + 1},
            )
        )
    return ParsedDocument(
        source_path=path,
        sections=sections,
        file_type="pptx",
        title=path.stem,
        metadata={"slide_count": len(prs.slides)},
    )
