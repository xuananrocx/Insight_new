from __future__ import annotations

import threading
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from src.core import accounts
from src.db import metadata_db as db
from src.knowledge import batch_upload as worker, ingestion
from src.knowledge.chunker import Chunk
from src.knowledge.parsers import xlsx_parser as xlsx


@pytest.fixture
def task(tmp_path, monkeypatch):
    db.init_db()
    monkeypatch.setattr(accounts, "enabled", False)
    path = tmp_path / "one.xlsx"
    path.write_bytes(b"placeholder")
    tid = worker.generate_task_id()
    db.create_upload_task(task_id=tid, kb_id="default", skip_mode="skip", auto_ingest=True,
                         files=[{"relative_path": path.name, "absolute_path": str(path), "file_size": 11}])
    return tid, db.list_upload_task_files(tid)[0]


def test_resume_processing_preserves_done_and_finishes(task, monkeypatch, tmp_path):
    tid, file = task
    db.append_upload_task_files(tid, [{"relative_path": "done.md", "absolute_path": str(tmp_path/'done.md'), "file_size": 1}])
    done = db.list_upload_task_files(tid)[1]
    db.update_upload_task_file_status(done["id"], "done")
    db.update_upload_task_file_status(file["id"], "processing")
    db.update_upload_task_progress(tid, status="paused", done=1)
    calls = []
    def ingest(path, **kwargs):
        calls.append(path)
        kwargs["progress"]("embedding", {"completed": 8, "total": 16, "detail": "8/16"})
        snapshot = db.get_upload_task(tid)
        assert snapshot["current_percent"] == 50
        assert snapshot["current_detail"] == "8/16"
        return {"ok": True}
    monkeypatch.setattr(ingestion, "ingest_single_file", ingest)
    monkeypatch.setattr(worker, "start_task", worker._process_task)
    assert worker.resume_task(tid)
    result = db.get_upload_task(tid)
    assert (result["status"], result["done"], len(calls)) == ("completed", 2, 1)
    assert result["current_percent"] is None and result["current_detail"] is None


def test_orphan_processing_cannot_remain_running(task):
    tid, file = task
    db.update_upload_task_file_status(file["id"], "processing")
    worker._process_task(tid)
    assert db.get_upload_task(tid)["status"] == "paused"


def test_empty_queue_reconciles_counts(task):
    tid, file = task
    db.update_upload_task_file_status(file["id"], "done")
    worker._process_task(tid)
    result = db.get_upload_task(tid)
    assert result["status"] == "completed" and result["done"] == 1


def test_duplicate_worker_is_not_started(task, monkeypatch):
    tid, _ = task
    started = []
    monkeypatch.setattr(threading, "Thread", lambda **kw: SimpleNamespace(start=lambda: started.append(kw)))
    try:
        worker.start_task(tid)
        worker.start_task(tid)
        assert not worker.resume_task(tid)
        assert len(started) == 1
    finally:
        worker._active_workers.discard(tid)


@pytest.mark.parametrize("status,expected", [("pending", True), ("failed", True), ("done", False)])
def test_skip_only_completed_records(tmp_path, monkeypatch, status, expected):
    db.init_db()
    path = tmp_path / "pending.md"
    path.write_text("text")
    fid = db.upsert_file(relative_path=str(path), absolute_path=str(path), content_hash="h",
                         file_size=4, file_type="md", kb_id="default")
    if status == "done": db.set_file_processed(fid, [])
    if status == "failed": db.set_file_failed(fid, "interrupted")
    calls = []
    monkeypatch.setattr(ingestion, "_process_one_document", lambda **kw: calls.append(kw))
    ingestion.ingest_single_file(path, skip_if_exists=True)
    assert bool(calls) == expected


def test_embedding_batches_and_cancel(monkeypatch):
    calls = []
    def embed(texts):
        calls.append(len(texts))
        return [[1.0] for _ in texts], "local"
    monkeypatch.setattr(ingestion.llm_client, "get_client", lambda: SimpleNamespace(embed=embed))
    chunks = [Chunk(str(i), {}) for i in range(70)]
    events = []
    vectors, _ = ingestion._embed_chunks(chunks, lambda s, p: events.append(p))
    assert calls == [32, 32, 6] and len(vectors) == 70
    assert events[-1]["completed"] == 70
    calls.clear()
    def cancel(stage, payload):
        if payload["completed"] >= 32: raise ingestion.TaskCancelledError()
    with pytest.raises(ingestion.TaskCancelledError): ingestion._embed_chunks(chunks, cancel)
    assert calls == [32]


def test_xlsx_sparse_bounds_types_and_cancel_cleanup(tmp_path, monkeypatch):
    path = tmp_path / "sparse.xlsx"
    wb = Workbook(); ws = wb.active; ws.title = "数据"
    ws.append(["名称", "数值", "日期"])
    ws.append(["中文|换行\n内容", 0, datetime(2026, 9, 22)])
    ws.append(["布尔", False, None])
    ws.cell(1048576, 16384).fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    wb.save(path); wb.close()
    sections = list(xlsx.iter_xlsx_sections(path))
    text = "\n".join(s.text for s in sections)
    assert len(text) < 500
    assert "中文\\|换行 内容" in text and "2026-09-22" in text
    assert "0" in text and "False" in text
    assert sections[0].extra["row_end"] == 3
    closed = []
    close = xlsx._DiskStrings.close
    def track(self):
        closed.append(True); close(self)
    monkeypatch.setattr(xlsx._DiskStrings, "close", track)
    def cancel(stage, payload): raise ingestion.TaskCancelledError()
    with pytest.raises(ingestion.TaskCancelledError): list(xlsx.iter_xlsx_sections(path, cancel))
    assert closed


def test_xlsx_blocks_keep_rows_and_release_on_close(tmp_path, monkeypatch):
    path = tmp_path / "large.xlsx"
    wb = Workbook(write_only=True); ws = wb.create_sheet()
    ws.append(["ID", "正文"])
    for i in range(2000): ws.append([f"row-{i}", "内容" * 30])
    wb.save(path); wb.close()
    sections = list(xlsx.iter_xlsx_sections(path))
    assert len(sections) > 2
    assert max(len(s.text) for s in sections) < xlsx.BLOCK_CHARS + 100
    text = "\n".join(s.text for s in sections)
    assert all(f"| row-{i} |" in text for i in range(2000))
    stream = xlsx.iter_xlsx_sections(path)
    next(stream); stream.close()
    # All workbook handles must be released, including early-close paths on Windows.
    path.rename(tmp_path / "closed.xlsx")


def test_shared_strings_rich_text_and_formula_cache(tmp_path):
    path = tmp_path / "shared.xlsx"
    wb = Workbook(); wb.active.append(["placeholder"]); wb.save(path); wb.close()
    with ZipFile(path) as archive: parts = {n: archive.read(n) for n in archive.namelist()}
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    parts["xl/sharedStrings.xml"] = f'<sst xmlns="{ns}"><si><r><t>中文</t></r><r><t>内容</t></r></si></sst>'.encode()
    parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(b'</Types>', b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')
    parts["xl/worksheets/sheet1.xml"] = f'<worksheet xmlns="{ns}"><sheetData><row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1"><f>1+1</f><v>2</v></c></row></sheetData></worksheet>'.encode()
    with ZipFile(path, 'w', ZIP_DEFLATED) as archive:
        for name, value in parts.items(): archive.writestr(name, value)
    assert "| 中文内容 | 2 |" in xlsx.parse_xlsx(path).full_text


@pytest.mark.asyncio
async def test_sse_snapshot_restores_and_unsubscribes(task):
    import json
    from src.api.routes_knowledge import stream_upload_task, _build_task_snapshot
    tid, _ = task
    db.update_upload_task_progress(tid, status="paused", current_file_path="one.xlsx",
                                  current_stage="embedding", current_percent=50, current_detail="8/16")
    assert _build_task_snapshot(tid)["id"] == tid
    assert _build_task_snapshot("missing") == {}
    response = await stream_upload_task(tid)
    events = [chunk async for chunk in response.body_iterator]
    snapshot = json.loads(events[0].split("data: ")[1])
    assert snapshot["current_percent"] == 50 and snapshot["current_detail"] == "8/16"
    assert "event: done" in events[-1]
    assert tid not in worker._subscribers


@pytest.mark.parametrize("fail_second", [False, True])
def test_ingestion_writes_bounded_batches_and_checkpoints(tmp_path, monkeypatch, fail_second):
    db.init_db()
    path = tmp_path / "batch.md"; path.write_text("document")
    chunks = [Chunk(f"text-{i}", {}) for i in range(70)]
    monkeypatch.setattr(ingestion, "chunk_document", lambda doc: chunks)
    monkeypatch.setattr(ingestion, "_is_ai_summary_enabled", lambda: False)
    monkeypatch.setattr(ingestion, "_embed_chunks", lambda batch: ([[1.0]*8 for _ in batch], "local"))
    batches = []
    def upsert(payloads, **kwargs):
        assert len(payloads) <= 32
        import json
        row = db.get_file_by_path(str(path), kb_id="default")
        from src.knowledge.ingest_transaction import pending_runs
        assert row["chunk_ids_json"] == "[]"
        assert any(len(json.loads(run["new_ids_json"])) == 70 for run in pending_runs())
        if fail_second and batches: raise RuntimeError("simulated write interruption")
        batches.append(len(payloads))
    monkeypatch.setattr(ingestion.vector_store, "upsert_chunks", upsert)
    monkeypatch.setattr(ingestion.vector_store, "get_or_create_collection", lambda name: SimpleNamespace(delete=lambda ids: None))
    monkeypatch.setattr(ingestion.bm25_index, "remove_chunks", lambda ids: None)
    def bm25(payloads):
        assert len(payloads) == 70
        assert all("embedding" not in payload for payload in payloads)
    monkeypatch.setattr(ingestion.bm25_index, "add_chunks", bm25)
    result = ingestion.ingest_single_file(path)
    assert result["ok"] == (not fail_second)
    assert batches == ([32] if fail_second else [32, 32, 6])
    row = db.get_file_by_path(str(path), kb_id="default")
    if fail_second:
        assert row is None  # Pending new-file record is cleaned, source is retained.
    else:
        assert row["status"] == "done" and row["chunk_ids_json"] != "[]"


def test_fast_batch_transitions_are_persisted(task, monkeypatch):
    tid, _ = task
    emitted = []
    monkeypatch.setattr(worker.time, "monotonic", lambda: 10.0)
    def ingest(path, **kwargs):
        report = kwargs["progress"]
        for done, operation in [(0, "computing"), (0, "writing"), (32, "batch_done"), (32, "computing")]:
            report("embedding", {"completed": done, "total": 286, "operation": operation, "detail": operation})
            snapshot = db.get_upload_task(tid)
            assert snapshot["current_detail"] == operation
            assert snapshot["current_percent"] == round(done / 286 * 100, 1)
        return {"ok": True}
    monkeypatch.setattr(ingestion, "ingest_single_file", ingest)
    worker._process_task(tid)
    assert db.get_upload_task(tid)["done"] == 1
