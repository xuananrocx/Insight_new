import json
import uuid
from types import SimpleNamespace

import pytest
from fastapi import HTTPException

from src.core import accounts, vector_store
from src.core.ingest_state import TaskCancelledError
from src.db import metadata_db as db
from src.knowledge import reading_index as ri, index_tasks, ingestion
from src.qa.knowledge_tools import KnowledgeTools


@pytest.fixture
def document(tmp_path,monkeypatch):
    db.init_db()
    monkeypatch.setattr(accounts,'enabled',False)
    kid=uuid.uuid4().hex
    db.create_kb(kid,'test',kid)
    path=tmp_path/'types.h'
    path.write_text('// header\n'*510+'struct Quote {\n'+''.join(f' int field_{i}; // description\n' for i in range(350))+' int visible_levels; // number of levels\n};\n','utf8')
    fid=db.upsert_file(path.name,str(path),ri.fingerprint(path),path.stat().st_size,'h',kb_id=kid)
    db.set_file_processed(fid,[])
    return db.get_file_by_path(path.name,kid)


def publish(file):
    builder=ri.build_file(file)
    with db.get_cursor() as cur:
        builder.publish(cur,file['id'],file['content_hash'])
    return builder


def test_manual_index_full_object_and_member_paging(document):
    publish(document)
    kb=KnowledgeTools(document['kb_id'])
    outline=kb.execute('get_document_outline',{'document_id':document['id'],'query':'Quote'})
    obj=outline['objects'][0]
    assert obj['boundary_known'] and obj['member_count']==351
    members=kb.execute('get_document_outline',{'document_id':document['id'],'object_id':obj['object_id'],'offset':350})
    assert 'visible_levels' in members['members'][0]['declaration']
    value=kb.execute('read_document',{'read_ref':obj['read_ref']})
    assert value['coverage']['object_complete']
    assert 'visible_levels' in value['evidence'][0]['text']
    assert value['evidence'][0]['text'].rstrip().endswith('};')
    # No vector lookups are needed to read a complete indexed object.
    assert kb.sources()[0]['reading']['object_complete']


def test_old_reference_rejected_after_rebuild_same_content(document):
    publish(document)
    kb=KnowledgeTools(document['kb_id'])
    ref=kb.outline(document['id'])['objects'][0]['read_ref']
    publish(document)
    assert kb.read_reference(ref)['error']=='document_changed'


def test_changed_source_and_cancel_preserve_published_index(document):
    from pathlib import Path
    first=publish(document)
    def cancel(): raise TaskCancelledError()
    with pytest.raises(TaskCancelledError): ri.build_file(document,cancel)
    Path(document['absolute_path']).write_text('changed','utf8')
    with pytest.raises(ValueError,match='已变化'): ri.build_file(document)
    assert ri.info(document)['artifact']==first.name


def test_index_tasks_persist_cancel_and_cross_kb_validation(document,monkeypatch):
    monkeypatch.setattr(index_tasks._pool,'submit',lambda *a:None)
    kid=document['kb_id']
    with pytest.raises(HTTPException): index_tasks.start(kid,[document['id']+99999])
    task=index_tasks.start(kid,[document['id']])
    with pytest.raises(HTTPException) as exc: index_tasks.start(kid,[document['id']])
    assert exc.value.status_code==409
    index_tasks.update(task['id'],cancel_requested=1)
    index_tasks.run(task['id'],kid)
    assert index_tasks.get(task['id'],kid)['state']=='cancelled'
    task=index_tasks.start(kid,[document['id']])
    index_tasks.run(task['id'],kid)
    assert index_tasks.get(task['id'],kid)['state']=='done'
    assert ri.info(document)


def test_import_builds_index_and_failed_update_keeps_old_version(document,monkeypatch):
    from pathlib import Path
    path=Path(document['absolute_path'])
    monkeypatch.setattr(ingestion,'_embed_chunks',lambda chunks:([[0.1,0.2]]*len(chunks),'test'))
    monkeypatch.setattr(ingestion,'_is_ai_summary_enabled',lambda:False)
    monkeypatch.setattr(vector_store,'upsert_chunks',lambda *a,**k:None)
    monkeypatch.setattr(vector_store,'get_or_create_collection',lambda *a,**k:SimpleNamespace(delete=lambda **kw:None))
    result=ingestion.ScanResult()
    ingestion._process_one_document(path,path.name,ri.fingerprint(path),path.stat().st_size,'h',None,document['kb_id'],result)
    current=db.get_file_by_path(path.name,document['kb_id'])
    assert result.updated==1 and ri.info(current)
    artifact=ri.info(current)['artifact']
    path.write_text('struct NewType { int replacement; };','utf8')
    def fail(*a,**k): raise ValueError('embedding failed')
    monkeypatch.setattr(ingestion,'_embed_chunks',fail)
    result=ingestion.ScanResult()
    ingestion._process_one_document(path,path.name,ri.fingerprint(path),path.stat().st_size,'h',None,document['kb_id'],result)
    current=db.get_file_by_path(path.name,document['kb_id'])
    assert result.failed==1 and current['content_hash']==document['content_hash']
    assert ri.info(current)['artifact']==artifact


def test_structure_does_not_balance_braces_in_comments_or_strings():
    text='struct A {\n const char* value="}"; // }\n int last;\n};\nstruct B { int b; };'
    objects=list(ri.structures(text))
    a=next(o for o in objects if o[1]=='A')
    assert a[4] and 'int last' in text[a[2]:a[3]]
    assert 'struct B' not in text[a[2]:a[3]]


def test_rebuild_retains_excel_cell_coordinates_and_merges(tmp_path,document):
    from openpyxl import Workbook
    path=tmp_path/'fields.xlsx'
    wb=Workbook(); ws=wb.active;ws.title='mapping'
    ws.append(['市场',None,'字段']);ws.merge_cells('A1:B1');ws.append(['沪市',None,'levels']);wb.save(path);wb.close()
    file={**document,'absolute_path':str(path),'content_hash':ri.fingerprint(path)}
    builder=ri.build_file(file)
    import sqlite3
    try:
        from contextlib import closing
        with closing(sqlite3.connect(builder.path)) as conn:
            extra=json.loads(conn.execute('select extra from sections').fetchone()[0])
        assert extra['merged_ranges']==['A1:B1']
        assert extra['cell_rows'][1]['values']=={'1':'沪市','3':'levels'}
        assert extra['cell_rows'][1]['row']==2
    finally: builder.discard()


def test_rebuild_uses_configured_feed_when_old_absolute_path_moved(document,monkeypatch):
    from pathlib import Path
    from src.core.config import settings
    actual=Path(document['absolute_path'])
    monkeypatch.setattr(type(settings),'feed_folder',property(lambda self: actual.parent))
    file={**document,'absolute_path':str(actual.parent/'missing.h'),'kb_id':'default'}
    builder=ri.build_file(file)
    builder.discard()


def test_index_endpoint_requires_management_before_starting_task(monkeypatch):
    from src.api import routes_indexes
    calls=[]
    def deny(*args):
        calls.append(args)
        raise HTTPException(403,'denied')
    monkeypatch.setattr(routes_indexes.permissions,'require_resource',deny)
    monkeypatch.setattr(index_tasks,'start',lambda *a:pytest.fail('unauthorized task started'))
    with pytest.raises(HTTPException):routes_indexes.rebuild('kb',routes_indexes.BuildRequest())
    assert calls==[('kb','kb','manage')]


def test_canonical_search_location_and_scope(document, monkeypatch):
    from pathlib import Path
    from src.qa import indexed_reading
    publish(document)
    kb=KnowledgeTools(document['kb_id'])
    source=Path(document['absolute_path']).read_text('utf8')
    excerpt=' int field_2; // description'
    location=indexed_reading.locate(kb,document,excerpt)
    assert location['status']=='exact'
    result=kb.read_reference(location['candidates'][0]['read_ref'])
    assert result['coverage']['object_complete']
    assert result['evidence'][0]['structure']['declaration_hint']['text']=='struct Quote {'
    assert result['read_status']['path']=='canonical_index'
    obj=kb.outline(document['id'],query='Quote')['objects'][0]
    chunks=[{'id':'inside','text':excerpt,'file_id':document['id']},
            {'id':'outside','text':'// header','file_id':document['id']}]
    selected, scope=indexed_reading.scope_chunks(kb,document,chunks,object_id=obj['object_id'])
    assert [c['id'] for c in selected]==['inside']
    assert selected[0]['_section_start']==source.index(excerpt)
    with pytest.raises(ValueError):indexed_reading.scope_chunks(kb,document,chunks,section=999)
    ambiguous=indexed_reading.locate(kb,document,'// header')
    assert ambiguous['status']=='ambiguous' and ambiguous['candidates_truncated']
    assert indexed_reading.locate(kb,document,'unfindable')['status']=='not_located'


def test_search_uses_canonical_reference_and_scopes_before_ranking(document,monkeypatch):
    publish(document)
    chunks=[{'id':'a','file_id':document['id'],'text':' int field_2; // description','section_index':0},
            {'id':'b','file_id':document['id'],'text':'// header','section_index':0}]
    db.set_file_processed(document['id'],['a','b'])
    kb=KnowledgeTools(document['kb_id'])
    monkeypatch.setattr(kb,'live_chunks',lambda ids:{c['id']:c for c in chunks if c['id'] in ids})
    monkeypatch.setattr(kb,'document_chunks',lambda f:(chunks,False))
    obj=kb.outline(document['id'],query='Quote')['objects'][0]
    res=kb.execute('search_document',{'document_id':document['id'],'query':'field_2','match_mode':'exact','object_id':obj['object_id']})
    assert res['status']=='found'
    ev=res['evidence'][0]
    assert ev['location']['status']=='exact_coordinate'
    assert kb.read_reference(ev['read_ref'])['coverage']['object_complete']
    outside=kb.execute('search_document',{'document_id':document['id'],'query':'header','match_mode':'exact','object_id':obj['object_id']})
    assert outside['status']=='no_match'
    assert res['search_scope']['selected_range']['object_id']==obj['object_id']


def test_excel_location_reads_original_coordinates(document,tmp_path):
    from openpyxl import Workbook
    from src.qa import indexed_reading
    path=tmp_path/'scope.xlsx'
    wb=Workbook();ws=wb.active;ws.title='Quotes'
    ws.append(['Title']);ws.merge_cells('A1:C1')
    ws.append(['Market','Field','Value']);ws.append(['X','levels',12])
    wb.create_sheet('Other').append(['different'])
    wb.save(path);wb.close()
    fid=db.upsert_file(path.name,str(path),ri.fingerprint(path),path.stat().st_size,'xlsx',kb_id=document['kb_id'])
    db.set_file_processed(fid,[])
    file=db.get_file_by_path(path.name,document['kb_id']);publish(file)
    kb=KnowledgeTools(file['kb_id'])
    hit=indexed_reading.locate(kb,file,'| X | levels | 12 |')
    assert hit['status']=='exact'
    assert hit['candidates'][0]['sheet']=='Quotes' and hit['candidates'][0]['row']==3
    res=kb.read_reference(hit['candidates'][0]['read_ref'])
    table=res['evidence'][0]['structure']['spreadsheet']
    assert table['rows'][0]['values']=={'1':'X','2':'levels','3':'12'}
    assert table['business_header_identified'] is False
    assert table['display_header']['row']==1
    assert res['previous_read_ref']
    chunks=[{'id':'a','text':'| X | levels | 12 |'},{'id':'b','text':'| different |'}]
    selected,_=indexed_reading.scope_chunks(kb,file,chunks,sheet='Quotes')
    assert [c['id'] for c in selected]==['a']


def test_tool_usage_descriptions_and_duplicate_boundaries():
    from src.qa.knowledge_tools import TOOLS
    specs={t['name']:t for t in TOOLS}
    assert not {'section','sheet','object_id'} & specs['search_document']['parameters']['properties'].keys()
    assert 'scope_ref' in specs['search_scope']['parameters']['properties']
    assert 'read_status' in specs['read_document']['description']
    assert 'reading_fallback' in specs['search_document']['description']
    a={'id':'a','file_id':1,'section_index':0,'_section_start':4,'text':'same'}
    b={**a,'id':'b'}
    c={**a,'id':'c','file_id':2}
    assert len(KnowledgeTools.unique_chunks([a,b,c]))==2
