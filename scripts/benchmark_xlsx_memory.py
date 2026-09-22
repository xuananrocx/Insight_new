import gc, tempfile, time, tracemalloc, json
from pathlib import Path
from openpyxl import Workbook, load_workbook
from src.knowledge.parsers.xlsx_parser import iter_xlsx_sections, _row_to_md
with tempfile.TemporaryDirectory() as temp:
    path = Path(temp)/"memory.xlsx"
    wb = Workbook(write_only=True); ws = wb.create_sheet()
    for i in range(6000): ws.append([str(i)+":"+str(j)+"中文内容"*50 for j in range(8)])
    wb.save(path); wb.close(); del wb, ws
    results = {}
    for mode in ["old", "streaming"]:
        gc.collect(); tracemalloc.start(); start = time.perf_counter()
        if mode == "old":
            wb = load_workbook(path, read_only=True, data_only=True)
            rows = list(wb.active.iter_rows(values_only=True))
            lines = [_row_to_md(list(row)) for row in rows]
            text = "\n".join(lines); chars = len(text)
            wb.close()
        else:
            chars = sum(len(section.text) for section in iter_xlsx_sections(path))
        results[mode] = {"peak_mib": round(tracemalloc.get_traced_memory()[1]/1024**2, 2), "seconds": round(time.perf_counter()-start, 2), "characters": chars}
        tracemalloc.stop()
        if mode == "old": del wb, rows, lines, text
    print(json.dumps(results))
